"""Genera el Excel de documentación E/S y causa-efecto a partir del análisis.

Uso: python generar_excel.py analisis.pkl config.json salida.xlsx
"""
import json
import pickle
import re
import sys
from collections import defaultdict, OrderedDict

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from glosario import traducir
from stl_logic import literals, size, RE_BIT, RE_BYTE, RE_TZ, bits_de, es_global

# ------------------------------------------------------------------ estilo
C_HEAD = '3B4B5C'
F_HEAD = PatternFill('solid', fgColor=C_HEAD)
F_SUB = PatternFill('solid', fgColor='D9DEE4')
F_RED = PatternFill('solid', fgColor='F4C7C3')
F_ORANGE = PatternFill('solid', fgColor='FCE4C4')
F_GRAY = PatternFill('solid', fgColor='E3E3E3')
F_A = PatternFill('solid', fgColor='C8E6C9')
F_P = PatternFill('solid', fgColor='D6E4F0')
F_B = PatternFill('solid', fgColor='F4C7C3')
FONT_HEAD = Font(bold=True, color='FFFFFF')
FONT_LINK = Font(color='1F4E9A', underline='single')
THIN = Side(style='thin', color='C9CED6')
BORDER = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top')
MAXTXT = 3000

D = None
SYM = {}
CFG = {}


def cut(s, n=MAXTXT):
    s = s or ''
    return s if len(s) <= n else s[:n - 40] + ' … (texto recortado; ver Referencias cruzadas)'


def sheet(wb, title, headers, widths, wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.fill = F_HEAD
        c.font = FONT_HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1] if i - 1 < len(widths) else 14
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'B2'
    ws._wrap_cols = set(wrap_cols)
    return ws


def finish(ws, ncols):
    ws.auto_filter.ref = f'A1:{get_column_letter(ncols)}{max(ws.max_row, 2)}'
    wc = ws._wrap_cols
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ncols):
        for c in row:
            c.alignment = WRAP if c.column in wc else TOP


# ------------------------------------------------------------------ descripción de operandos
def addr_key(op):
    m = RE_BIT.match(op or '')
    if m:
        return (m.group(1), int(m.group(2)), int(m.group(3)))
    m = RE_BYTE.match(op or '')
    if m:
        return (m.group(1) + m.group(2), int(m.group(3)), 0)
    m = RE_TZ.match(op or '')
    if m:
        return (m.group(1), int(m.group(2)), 0)
    return ('~' + (op or ''), 0, 0)


def db_sym(dbn):
    s = SYM.get(dbn)
    return s['Symbol'] if s else dbn


def nombre(op):
    """Operando con símbolo, p. ej. E125.1 «325B125.1»."""
    if op is None:
        return '?'
    if op.startswith('~'):
        blk, _, var = op[1:].partition('.')
        return f'#{var} (temporal {blk})'
    if op.startswith('?'):
        return f'{op[1:]} (parámetro sin conectar)'
    s = SYM.get(op)
    if s:
        return f'{op} «{s["Symbol"]}»'
    m = re.match(r'^(DB\d+)\.(.*)$', op)
    if m:
        return f'{m.group(1)}.{m.group(2)} «{db_sym(m.group(1))}»' if not m.group(2).startswith('DB') else \
            f'{op} «{db_sym(m.group(1))}»'
    return op


def descr(op):
    """Texto legible para la forma en lenguaje claro."""
    if op is None:
        return '?'
    if op.startswith('~'):
        blk, _, var = op[1:].partition('.')
        return f'variable temporal #{var} de {blk}'
    s = SYM.get(op)
    if s:
        com = s['Comment'].strip()
        if com and com.upper() not in ('RESERVE', '.', 'RESERVA'):
            return f'{traducir(com)} ({op})'
        return f'{s["Symbol"]} ({op})'
    info = D['sympath_info'].get(op)
    m = re.match(r'^(DB\d+)\.(.*)$', op)
    if m:
        base = f'"{db_sym(m.group(1))}".{m.group(2)}'
        if info and info[2]:
            return f'{traducir(info[2])} ({base})'
        return base
    return op


def render(e, txt=False, top=True):
    t = e[0]
    j_and = ' Y ' if txt else ' ∧ '
    j_or = ' O ' if txt else ' ∨ '
    neg = 'NO ' if txt else '¬'
    if t == 'c':
        return ('siempre' if e[1] else 'nunca') if txt else ('1' if e[1] else '0')
    if t == 'l':
        name = descr(e[1]) if txt else nombre(e[1])
        if txt and RE_TZ.match(e[1] or ''):
            name = 'temporizador ' + name if e[1].startswith('T') else 'contador ' + name
        return (neg if e[2] else '') + name
    if t == 'a':
        return j_and.join(render(x, txt, False) if x[0] != 'o' else '(' + render(x, txt, False) + ')' for x in e[1])
    if t == 'o':
        s = j_or.join(render(x, txt, False) for x in e[1])
        return s
    if t == 'n':
        return neg + '(' + render(e[1], txt, False) + ')'
    if t == 'x':
        return '(' + render(e[1], txt, False) + (' XOR ' if txt else ' ⊕ ') + render(e[2], txt, False) + ')'
    if t == 'e':
        f = 'flanco positivo' if e[1] == 'P' else 'flanco negativo'
        return f'{f}[{render(e[2], txt, False)}] (memoria {e[3]})' if txt else \
            f'{"FP" if e[1] == "P" else "FN"}({render(e[2], txt, False)}; {e[3]})'
    if t == 'k':
        return f'[{e[1]} {e[2]} {e[3]}]'
    if t == 'q':
        return f'⟨{e[1]}⟩'
    if t == 't':
        return f'valor transferido desde {e[1]}' if txt else f'T ← {e[1]}'
    return str(e)


def ubicacion(w):
    b = w['block']
    bs = D['blocks'].get(b, {}).get('Symbol') or ''
    s = f'{b}' + (f' «{bs}»' if bs else '') + f' NW{w["nw"]}'
    if w.get('chain'):
        via = ' ← '.join(f'{c[0]} NW{c[1]}' for c in reversed(w['chain']))
        s += f' (vía {via})'
    if w.get('prefix'):
        s += f' [instancia {w["prefix"].rstrip(".")}]'
    if w.get('uncalled'):
        s += ' [BLOQUE NO LLAMADO]'
    return s


KIND_TXT = {'assign': '=', 'set': 'S', 'reset': 'R', 'edge': 'FP/FN', 'transfer': 'T', 'call_out': 'CALL (salida)',
            'timer_start': 'arranque temporizador', 'count_up': 'ZV', 'count_down': 'ZR'}


def cond_text(w, txt):
    ex = w['expr']
    pc = w['pc']
    if ex[0] == 't':
        s = render(ex, txt)
    else:
        s = render(ex, txt)
    if pc != ('c', True):
        s += ('  — SOLO SI: ' if txt else '  | camino: ') + render(pc, txt)
    return s


# ------------------------------------------------------------------ BMK / plano
def bmk_de(op):
    """(hoja, clase, bmk, coherente) deducidos del símbolo o del comentario."""
    s = SYM.get(op)
    m = RE_BIT.match(op or '')
    if not s or not m:
        return None
    addr = f'{m.group(2)}.{m.group(3)}'
    for src in (s['Symbol'], s['Comment']):
        for t in re.findall(r'(\d{3})\s?(KM|[A-Z])\s?(\d+\.\d)', src or ''):
            if t[2] == addr:
                exp = (200 if m.group(1) == 'E' else 400) + int(m.group(2))
                return (int(t[0]), t[1], f'{t[0]}{t[1]}{t[2]}', int(t[0]) == exp)
    return None


CLASE = {'B': 'Sensor / detector', 'S': 'Pulsador / interruptor / selector', 'KM': 'Contactor de motor',
         'K': 'Relé / contactor', 'Y': 'Electroválvula', 'H': 'Piloto / señalización', 'Q': 'Interruptor / guardamotor',
         'V': 'Indicador / LED', 'M': 'Motor', 'P': 'Señalización acústica / indicador', 'U': 'Convertidor / módulo'}


def elemento_campo(op):
    s = SYM.get(op)
    txt = ((s['Symbol'] + ' ' + s['Comment']) if s else '').upper()
    bm = bmk_de(op)
    kw = [('LICHTSCHRANKE', 'Fotocélula'), ('LICHTTASTER', 'Fotocélula de detección directa'),
          ('INITIATOR', 'Detector inductivo'), (' INI', 'Detector inductivo'), ('REED', 'Contacto reed'),
          ('NOT-AUS', 'Seta de emergencia / circuito de emergencia'), ('ESTOP', 'Seta de emergencia'),
          ('TASTER', 'Pulsador'), ('WAHLSCHALTER', 'Selector'), ('SCHALTER', 'Interruptor'),
          ('MOTORSCHUTZ', 'Contacto auxiliar de guardamotor'), ('ANTR', 'Contactor de accionamiento (motor)'),
          ('VENTIL', 'Electroválvula'), ('LEUCHTMELDER', 'Piloto luminoso'), ('LM ', 'Piloto luminoso'),
          ('HUPE', 'Bocina'), ('HORN', 'Bocina'), ('SICHERH', 'Barrera / dispositivo de seguridad'),
          ('TUER', 'Interruptor de puerta'), ('TÜR', 'Interruptor de puerta'), ('WAAGE', 'Báscula'),
          ('LED', 'Indicador LED')]
    for k, v in kw:
        if k in ' ' + txt:
            return v + (f' (clase BMK {bm[1]})' if bm else '')
    if bm:
        return CLASE.get(bm[1], f'Clase BMK {bm[1]}') + ' (por clase BMK)'
    return 'No localizado'


def posicion_mecanica(op):
    s = SYM.get(op)
    if not s:
        return ''
    t = s['Comment'] or ''
    parts = []
    m = re.search(r'\b(\d{1,3}M\d)\b', t)
    if m:
        parts.append(f'accionamiento {m.group(1)}')
    m = re.search(r'POS(?:ITION|\.)?\s*(\d+)', t, re.I)
    if m:
        parts.append(f'posición mecánica {m.group(1)}')
    m = re.search(r'\b(PAH|PZM|SPF|FS|KA|L)\s?(\d+)\b', t)
    if m:
        parts.append(f'grupo {m.group(1)} {m.group(2)}')
    return ', '.join(parts)


def link_plano(ws, row, col, hoja):
    if hoja is None:
        ws.cell(row=row, column=col, value='No localizado')
        return
    f = (f'=HYPERLINK(Portada!$C${CFG["_cell_pdf"]}&"#page="&({hoja}+Portada!$C${CFG["_cell_off"]}),'
         f'"Hoja {hoja}")')
    c = ws.cell(row=row, column=col, value=f)
    c.font = FONT_LINK


# ------------------------------------------------------------------ principal
def main(pkl, cfgpath, dst):
    global D, SYM, CFG
    D = pickle.load(open(pkl, 'rb'))
    CFG = json.load(open(cfgpath, encoding='utf-8'))
    CFG['_cell_pdf'], CFG['_cell_off'] = 16, 17     # filas fijas en Portada (ver rows0)
    SYM = D['sym']
    blocks = D['blocks']
    W, RD, TM, XR = D['W'], D['RD'], D['TM'], D['xref']

    # -------- índices
    writes = defaultdict(list)
    for w in W:
        writes[w['op']].append(w)
        for b in bits_de(w['op']):
            writes[b].append(dict(w, via_byte=w['op']))
    reads_by = defaultdict(list)
    for r in RD:
        reads_by[r['op']].append(r)
        for b in bits_de(r['op']):
            reads_by[b].append(dict(r, via_byte=r['op']))
    xref_by = defaultdict(list)
    for x in XR:
        xref_by[x[0]].append(x)
        for b in bits_de(x[0]):
            xref_by[b].append((b,) + x[1:4] + (x[4] + f' (vía {x[0]})',))

    # dependencias: escritura -> literales (con polaridad / rol)
    deps = defaultdict(lambda: defaultdict(set))    # destino -> origen -> {'A','P','B'}
    for w in W:
        if w['expr'][0] == 't' or w['kind'] in ('call_out',):
            for o in w['expr'][2] if w['expr'][0] == 't' else ():
                if isinstance(o, str):
                    deps[w['op']][o].add('P')
            continue
        kind = w['kind']
        ex = w['expr']
        terms = ex[1] if ex[0] == 'o' else (ex,)
        for t in terms:
            fs = t[1] if t[0] == 'a' else (t,)
            for i, f in enumerate(fs):
                for o, n in literals(f):
                    if kind == 'reset':
                        deps[w['op']][o].add('B' if not n else 'P')
                    elif n:
                        deps[w['op']][o].add('B')
                    else:
                        deps[w['op']][o].add('A' if (i == 0 and f[0] == 'l') else 'P')
        for o, n in literals(w['pc']):
            deps[w['op']][o].add('B' if n else 'P')

    outputs_written = sorted({o for o in writes if RE_BIT.match(o) and o.startswith('A')}, key=addr_key)
    inv = defaultdict(lambda: defaultdict(set))     # origen -> destino -> roles
    for dst_, srcs in deps.items():
        for src, roles in srcs.items():
            inv[src][dst_] |= roles

    def alcanzables(src, maxd=5):
        """Salidas A alcanzables desde src a través de marcas/DB/temporizadores."""
        res = {}
        frontier = [(src, 0)]
        seen = {src}
        while frontier:
            n, d = frontier.pop(0)
            for t, roles in inv.get(n, {}).items():
                if t in seen:
                    continue
                seen.add(t)
                if t.startswith('A') and RE_BIT.match(t):
                    res.setdefault(t, (d + 1, roles))
                if d + 1 < maxd and not (t.startswith('A') and RE_BIT.match(t)):
                    frontier.append((t, d + 1))
        return res

    wb = Workbook()
    wb.remove(wb.active)

    # ===================================================== Portada (se rellena al final)
    ws0 = wb.create_sheet('Portada')

    # ===================================================== Bloques
    callers = defaultdict(set)
    for c in D['CALLS']:
        if c['callee']:
            callers[c['callee']].add(f'{c["caller"]} NW{c["nw"]}')
    calls_of = defaultdict(set)
    for c in D['CALLS']:
        if c['callee']:
            calls_of[c['caller']].add(c['callee'] + (f', {c["inst"]}' if c['inst'] else ''))
    wsb = sheet(wb, 'Bloques', ['Bloque', 'Símbolo', 'Tipo', 'Lenguaje', 'Título / comentario', 'Familia', 'Autor',
                                'Versión', 'Último cambio código', 'Segmentos', 'Llamado desde', 'Llama a',
                                'Instancia de', 'Observaciones'],
                [10, 26, 7, 9, 48, 12, 12, 8, 17, 10, 40, 40, 12, 40], wrap_cols=(5, 11, 12, 14))
    order = {'OB': 0, 'FB': 1, 'FC': 2, 'DB': 3, 'UDT': 4, 'SFB': 5, 'SFC': 6, 'VAT': 7}
    lang_es = {'AWL': 'AWL/STL', 'FUP': 'FUP/FBD', 'KOP': 'KOP/LAD', 'SCL': 'SCL', 'GRAPH': 'GRAPH', 'unkown': '—'}
    blist = sorted(blocks.values(), key=lambda b: (order.get(b['Type'], 9), int(re.sub(r'\D', '', b['Name']) or 0)))
    for b in blist:
        n = b['Name']
        obs = []
        if n in D['no_llamados']:
            obs.append('No se llama desde ningún OB (análisis estático): código sin ejecutar o llamado de forma dinámica')
        if b.get('KnowHow'):
            obs.append('Bloque con protección know-how')
        if D['tpl'].get(n, {}).get('indirect'):
            obs.append('Contiene direccionamiento indirecto / listas de salto: referencias cruzadas posiblemente incompletas')
        inst_of = f'FB{b["FB"]}' if b['Type'] == 'DB' and b.get('IsInstance') else ''
        wsb.append([n, b.get('Symbol') or '', b['Type'], lang_es.get(b.get('Lang'), b.get('Lang')),
                    b.get('Title') or b.get('SymComment') or '', b.get('Family') or '', b.get('Author') or '',
                    b.get('Version') or '', b.get('LastCodeChange') or '', len(b.get('Networks') or []) or '',
                    ', '.join(sorted(callers.get(n, []))), ', '.join(sorted(calls_of.get(n, []))), inst_of,
                    '; '.join(obs)])
    finish(wsb, 14)

    # ===================================================== Referencias cruzadas (primero, para anclas)
    wsx = sheet(wb, 'Referencias cruzadas', ['Variable', 'Símbolo', 'Bloque', 'Símbolo bloque', 'Segmento (NW)',
                                             'Lectura/Escritura', 'Instrucción'],
                [16, 28, 9, 26, 10, 12, 60], wrap_cols=(7,))
    anchor = {}
    xr_sorted = sorted(set(XR), key=lambda x: (addr_key(x[0]), x[1], x[2], x[3]))
    for x in xr_sorted:
        op = x[0]
        if op not in anchor:
            anchor[op] = wsx.max_row + 1
        s = SYM.get(op)
        wsx.append([op, s['Symbol'] if s else '', x[1], blocks.get(x[1], {}).get('Symbol') or '', x[2],
                    {'R': 'Lectura', 'W': 'Escritura'}[x[3]], x[4][:300]])
    # anclas para bits accedidos por byte/palabra
    for op in list(xref_by):
        if op not in anchor:
            for x in xref_by[op]:
                src = re.search(r'\(vía (\S+)\)$', x[4])
                if src and src.group(1) in anchor:
                    anchor[op] = anchor[src.group(1)]
                    break
    finish(wsx, 7)

    def link_addr(ws, row, col, op):
        c = ws.cell(row=row, column=col, value=op)
        if op in anchor:
            c.hyperlink = f"#'Referencias cruzadas'!A{anchor[op]}"
            c.font = FONT_LINK

    def bloques_uso(op):
        xs = xref_by.get(op, [])
        g = OrderedDict()
        for x in sorted(xs, key=lambda x: (x[1], x[2])):
            g.setdefault(x[1], set()).add((x[2], x[3]))
        out = []
        for b, nws in g.items():
            nn = ', '.join(f'NW{n}{"(E)" if rw == "W" else ""}' for n, rw in sorted(nws))
            out.append(f'{b}: {nn}')
        return '; '.join(out)

    # ===================================================== avisos
    avisos = []   # (severidad, tipo, variable, detalle, ubicación)

    def aviso(sev, tipo, var, det, ubi=''):
        if tipo == 'Usada sin declarar' and RE_BIT.match(var or ''):
            m_ = RE_BIT.match(var)
            n_ = int(m_.group(2))
            cub = [f'{m_.group(1)}B{n_}', f'{m_.group(1)}W{n_}', f'{m_.group(1)}W{n_ - 1}', f'{m_.group(1)}D{n_}',
                   f'{m_.group(1)}D{n_ - 1}', f'{m_.group(1)}D{n_ - 2}', f'{m_.group(1)}D{n_ - 3}']
            if any(c_ in SYM for c_ in cub):
                return
            if var.startswith('M'):
                sev = 'Info'
        avisos.append((sev, tipo, var, det, ubi))

    # ===================================================== Entradas
    used_ops = set(xref_by) | set(writes) | set(reads_by)
    ins = sorted({o for o in SYM if o.startswith('E') and RE_BIT.match(o)} |
                 {o for o in used_ops if o.startswith('E') and RE_BIT.match(o)}, key=addr_key)
    wse = sheet(wb, 'Entradas', ['Dirección', 'Símbolo', 'Descripción programa', 'Descripción (ES, glosario)',
                                 'Descripción plano', 'Tipo de señal', 'Elemento de campo', 'BMK', 'Módulo/Canal',
                                 'Borne', 'Página plano', 'Función mecánica (del comentario)', 'Qué activa / permite',
                                 'Qué bloquea', 'Bloques donde se usa', 'Observaciones'],
                [10, 18, 32, 32, 16, 22, 24, 13, 16, 10, 11, 24, 60, 50, 40, 50],
                wrap_cols=(3, 4, 6, 7, 12, 13, 14, 15, 16))
    stats = {'E': 0, 'E_plano': 0, 'A': 0, 'A_plano': 0}
    inv_writes_E = {}
    for o in ins:
        s = SYM.get(o, {})
        stats['E'] += 1
        bm = bmk_de(o)
        obs = []
        tipo = 'No determinable sin plano'
        ws_e = [w for w in writes.get(o, []) if not w.get('via_byte')]
        if ws_e:
            locs = '; '.join(sorted({ubicacion(w) for w in ws_e}))
            inv_e = any(w['kind'] == 'assign' and w['expr'] == ('l', o, True) for w in ws_e)
            if inv_e:
                tipo = 'Probable NC / activa a nivel bajo (el programa la invierte)'
            obs.append(f'AVISO: el programa escribe esta entrada ({locs})')
            aviso('Aviso', 'Escritura sobre entrada', o,
                  ('Inversión de la señal en la imagen de proceso (U NOT / =)' if inv_e else
                   'La entrada se fuerza/escribe por programa: ' + ', '.join(sorted({KIND_TXT.get(w["kind"], w["kind"]) for w in ws_e}))),
                  locs)
        com = (s.get('Comment') or '').upper()
        if 'NOT-AUS' in com or 'SICHERH' in com or 'NOT-HALT' in com:
            if tipo.startswith('No'):
                tipo = 'NC probable (circuito de seguridad) — verificar en plano'
        hoja = None
        if bm:
            hoja = bm[0]
            stats['E_plano'] += 1
            if not bm[3]:
                obs.append(f'DISCREPANCIA: el BMK {bm[2]} indica hoja {bm[0]}, que no sigue la numeración 200+byte del resto de entradas')
                aviso('Discrepancia', 'Hoja/BMK incoherente', o, f'BMK {bm[2]} → hoja {bm[0]}; por convención sería {200 + addr_key(o)[1]}', '')
        # activa / bloquea (directo e indirecto)
        direct = inv.get(o, {})
        act, blq = [], []
        for t, roles in sorted(direct.items(), key=lambda t: addr_key(t[0])):
            if not (t.startswith('A') or t.startswith('M')):
                continue
            tag = nombre(t)
            if roles & {'A', 'P'}:
                act.append(f'{tag} ({"/".join(sorted(roles & {"A", "P"}))})')
            if 'B' in roles:
                blq.append(tag)
        indirect = alcanzables(o)
        ind_txt = [f'{nombre(t)} [indirecto, {d} niveles]' for t, (d, r) in
                   sorted(indirect.items(), key=lambda t: addr_key(t[0])) if t not in direct and d > 1]
        act_s = '\n'.join(act[:25]) + (f'\n… y {len(act) - 25} más' if len(act) > 25 else '')
        if ind_txt:
            act_s += ('\n' if act_s else '') + 'Indirectamente (a través de marcas / DB / temporizadores):\n' + \
                '\n'.join(ind_txt[:15]) + (f'\n… y {len(ind_txt) - 15} más' if len(ind_txt) > 15 else '')
        blq_s = '\n'.join(blq[:25]) + (f'\n… y {len(blq) - 25} más' if len(blq) > 25 else '')
        uso = bloques_uso(o)
        if not uso:
            obs.append('Sin uso en el programa')
            if s:
                aviso('Sin uso', 'Entrada declarada sin uso', o, f'{s.get("Symbol")} — {s.get("Comment")}')
        else:
            if com.strip() in ('RESERVE', 'RESERVA', 'RESERVE.') or (s.get('Symbol') or '').upper().startswith('RESERVE'):
                obs.append('DISCREPANCIA: el símbolo/comentario dice RESERVA pero la entrada se usa en el programa')
                aviso('Discrepancia', 'Reserva usada', o, f'Comentario "{s.get("Comment")}" pero se usa en: {uso}', uso)
        if not s:
            obs.append('AVISO: usada sin declarar en la tabla de símbolos')
            aviso('Aviso', 'Usada sin declarar', o, 'Operando sin entrada en la tabla de símbolos', uso)
        obs.append('Plano no disponible: descripción, módulo, borne y cable sin verificar')
        r = wse.max_row + 1
        wse.append([None, s.get('Symbol', ''), s.get('Comment', ''), traducir(s.get('Comment', '')),
                    'No localizado', tipo, elemento_campo(o), bm[2] if bm else 'No localizado',
                    'No localizado', 'No localizado', None, posicion_mecanica(o), cut(act_s), cut(blq_s), uso,
                    '; '.join(obs)])
        link_addr(wse, r, 1, o)
        link_plano(wse, r, 11, hoja)
    finish(wse, 16)

    # ===================================================== Salidas
    outs = sorted({o for o in SYM if o.startswith('A') and RE_BIT.match(o)} |
                  {o for o in used_ops if o.startswith('A') and RE_BIT.match(o)}, key=addr_key)
    wss = sheet(wb, 'Salidas', ['Dirección', 'Símbolo', 'Descripción', 'Descripción (ES, glosario)', 'Actuador', 'BMK',
                                'Módulo/Canal', 'Borne', 'Página plano', 'Función mecánica (del comentario)',
                                'Condiciones de activación (lenguaje claro)', 'Condiciones de activación (booleana)',
                                'Condiciones de desactivación', 'Enclavamientos / seguridades',
                                'Temporizadores y contadores implicados', 'Bloque/Segmento de escritura',
                                'Bloques donde se lee', 'Observaciones'],
                [10, 18, 30, 30, 24, 13, 14, 10, 11, 24, 70, 60, 55, 45, 30, 50, 35, 50],
                wrap_cols=(3, 4, 5, 10, 11, 12, 13, 14, 15, 16, 17, 18))
    SAFE_KW = ('NOT-AUS', 'NOT-HALT', 'SICHERH', 'MOTORSCHUTZ', 'TUER', 'TÜR', 'PNOZ', 'ESTOP', 'STÖR', 'STOER',
               'SCHUTZ', 'MS_', 'MMS', 'EMERG', 'MUTING', 'ENDLAGE')
    double_assign = {}
    for o in outs:
        s = SYM.get(o, {})
        stats['A'] += 1
        bm = bmk_de(o)
        ws_o = writes.get(o, [])
        obs = []
        act_t, act_b, des, encl, tms, locs = [], [], [], [], set(), []
        assigns = [w for w in ws_o if w['kind'] == 'assign' and not w.get('via_byte')]
        sets = [w for w in ws_o if w['kind'] == 'set']
        resets = [w for w in ws_o if w['kind'] == 'reset']
        others = [w for w in ws_o if w['kind'] not in ('assign', 'set', 'reset') or w.get('via_byte')]
        for w in assigns + sets:
            pref = '= ' if w['kind'] == 'assign' else 'S '
            act_t.append(f'• [{pref.strip()}] {cond_text(w, True)}\n   ({ubicacion(w)})')
            act_b.append(f'{o} {"=" if w["kind"] == "assign" else ":S"} {cond_text(w, False)}')
        for w in resets:
            des.append(f'• [R] {cond_text(w, True)}\n   ({ubicacion(w)})')
        if assigns:
            des.insert(0, '• Asignación (=): se desactiva en cuanto deja de cumplirse la condición de activación'
                       + (' (el segmento solo se ejecuta si se cumple la condición de camino indicada)'
                          if any(w['pc'] != ('c', True) for w in assigns) else ''))
        for w in others:
            src = w.get('via_byte') or o
            act_t.append(f'• [{KIND_TXT.get(w["kind"], w["kind"])}{" " + src if w.get("via_byte") else ""}] '
                         f'{render(w["expr"], True)}\n   ({ubicacion(w)})')
            act_b.append(f'{src} {KIND_TXT.get(w["kind"], w["kind"])} {render(w["expr"], False)}')
        for w in ws_o:
            locs.append(ubicacion(w))
            for lo, n in literals(w['expr']) + literals(w['pc']) if w['expr'][0] != 't' else []:
                if RE_TZ.match(lo):
                    tms.add(lo)
                ss = SYM.get(lo, {})
                txt = ((ss.get('Symbol') or '') + ' ' + (ss.get('Comment') or '')).upper()
                if any(k in txt for k in SAFE_KW) or (n and w['kind'] in ('assign', 'set')) or \
                        (not n and w['kind'] == 'reset'):
                    e = f'{"NO " if n and w["kind"] != "reset" else ""}{descr(lo)}' if w['kind'] != 'reset' else \
                        f'{"NO " if n else ""}{descr(lo)} → desactiva'
                    if e not in encl:
                        encl.append(e)
        # temporizadores con preselección
        tm_txt = []
        for t in sorted(tms, key=addr_key):
            pres = sorted({str(x['preset']) for x in D['TM'] if x['op'] == t and x['preset']})
            tps = sorted({x['type'] for x in D['TM'] if x['op'] == t})
            tm_txt.append(f'{nombre(t)}' + (f' {"/".join(tps)} {", ".join(pres)}' if tps else ''))
        # avisos
        sites = {(w['block'], w['nw'], w.get('prefix')) for w in assigns}
        if len(sites) > 1:
            obs.append(f'AVISO: doble asignación (= en {len(sites)} sitios)')
            excl = all(w['pc'] != ('c', True) for w in assigns)
            aviso('Aviso', 'Doble asignación', o, f'Se asigna con "=" en {len(sites)} lugares; prevalece el último ejecutado en el ciclo'
                  + ('. Están en ramas con condición de salto distinta (manual/automático…): comprobar que son excluyentes' if excl else ''),
                  '; '.join(sorted({ubicacion(w) for w in assigns})))
        if assigns and sets:
            obs.append('AVISO: mezcla de asignación (=) y SET')
            aviso('Aviso', 'Mezcla = con S/R', o, 'La salida se escribe con "=" y también con SET: el resultado depende del orden de ejecución',
                  '; '.join(sorted({ubicacion(w) for w in assigns + sets + resets}))[:900])
        elif assigns and resets:
            obs.append('Reset previo + asignación (=): patrón habitual de puesta a cero, revisar orden')
            aviso('Info', 'Reset previo + asignación', o, 'Se resetea (R) y se asigna (=) en otro sitio; patrón Langhammer de puesta a cero de salidas cuando no hay automático',
                  '; '.join(sorted({ubicacion(w) for w in assigns + resets}))[:900])
        if any(w.get('uncalled') for w in ws_o):
            obs.append('AVISO: se escribe en un bloque que no se llama')
        if not ws_o:
            if s:
                obs.append('Sin escritura en el programa' + (' ni lectura' if not bloques_uso(o) else ''))
                aviso('Sin uso', 'Salida declarada sin escritura', o, f'{s.get("Symbol")} — {s.get("Comment")}',
                      bloques_uso(o))
        if not s:
            obs.append('AVISO: usada sin declarar en la tabla de símbolos')
            aviso('Aviso', 'Usada sin declarar', o, 'Operando sin entrada en la tabla de símbolos', bloques_uso(o))
        com = (s.get('Comment') or '').upper().strip()
        if ws_o and (com in ('RESERVE', 'RESERVA') or (s.get('Symbol') or '').upper().startswith('RESERVE')):
            obs.append('DISCREPANCIA: comentario RESERVA pero la salida se escribe')
            aviso('Discrepancia', 'Reserva usada', o, f'Comentario "{s.get("Comment")}" pero se escribe en el programa',
                  '; '.join(sorted(set(locs)))[:600])
        hoja = None
        if bm:
            hoja = bm[0]
            stats['A_plano'] += 1
            if not bm[3]:
                obs.append(f'DISCREPANCIA: BMK {bm[2]} → hoja {bm[0]}, fuera de la numeración 400+byte')
                aviso('Discrepancia', 'Hoja/BMK incoherente', o, f'BMK {bm[2]} → hoja {bm[0]}; por convención sería {400 + addr_key(o)[1]}')
        if any(w['pc'][0] == 'q' and 'compleja' in w['pc'][1] for w in ws_o):
            obs.append('Lógica con saltos complejos: condición resumida')
        obs.append('Plano no disponible: actuador, módulo y borne sin verificar')
        rd = sorted({f'{r["block"]} NW{r["nw"]}' for r in reads_by.get(o, [])})
        r = wss.max_row + 1
        wss.append([None, s.get('Symbol', ''), s.get('Comment', ''), traducir(s.get('Comment', '')), elemento_campo(o),
                    bm[2] if bm else 'No localizado', 'No localizado', 'No localizado', None, posicion_mecanica(o),
                    cut('\n'.join(act_t)) or 'No se escribe en el programa', cut('\n'.join(act_b)), cut('\n'.join(des)),
                    cut('\n'.join(encl[:40])), '\n'.join(tm_txt), cut('\n'.join(sorted(set(locs)))),
                    cut(', '.join(rd[:60]) + (f' … (+{len(rd) - 60})' if len(rd) > 60 else '')), '; '.join(obs)])
        link_addr(wss, r, 1, o)
        link_plano(wss, r, 9, hoja)
    finish(wss, 18)

    # ===================================================== Marcas
    marks = sorted({o for o in SYM if o.startswith('M')} | {o for o in used_ops if re.match(r'^M[BWD]?\d', o)},
                   key=addr_key)
    wsm = sheet(wb, 'Marcas', ['Dirección', 'Símbolo', 'Tipo', 'Función (comentario)', 'Función (ES, glosario)',
                               'Cómo se activa', 'Cómo se desactiva', 'Qué activa / permite', 'Remanente',
                               'Bloques de escritura', 'Bloques de lectura', 'Observaciones'],
                [10, 22, 7, 30, 30, 70, 50, 50, 22, 40, 40, 40], wrap_cols=(4, 5, 6, 7, 8, 10, 11, 12))
    for o in marks:
        s = SYM.get(o, {})
        ws_o = writes.get(o, [])
        act, des, obs = [], [], []
        for w in ws_o:
            if w['kind'] == 'reset':
                des.append(f'• [R] {cond_text(w, True)} ({ubicacion(w)})')
            else:
                k = KIND_TXT.get(w['kind'], w['kind'])
                act.append(f'• [{k}{" " + w["via_byte"] if w.get("via_byte") else ""}] {cond_text(w, True)} ({ubicacion(w)})')
                if w['kind'] == 'assign':
                    des.append('• Asignación (=): se desactiva al dejar de cumplirse la condición')
        des = list(OrderedDict.fromkeys(des))
        tgt = []
        for t, roles in sorted(inv.get(o, {}).items(), key=lambda t: addr_key(t[0])):
            if t.startswith('~'):
                continue
            tgt.append(f'{nombre(t)} ({"/".join(sorted(roles))})')
        k = addr_key(o)
        rem = ('Sí, si se mantiene el ajuste por defecto (MB0–MB15)' if k[1] <= 15 else
               'No, con el ajuste por defecto (MB0–MB15)') + ' — sin verificar (HW Config no exportado)'
        sites = {(w['block'], w['nw'], w.get('prefix')) for w in ws_o if w['kind'] == 'assign' and not w.get('via_byte')}
        if len(sites) > 1:
            obs.append(f'AVISO: doble asignación (= en {len(sites)} sitios)')
            aviso('Aviso', 'Doble asignación', o, f'Marca asignada con "=" en {len(sites)} lugares',
                  '; '.join(sorted({ubicacion(w) for w in ws_o if w['kind'] == 'assign'}))[:900])
        uso = bloques_uso(o)
        if not uso and not ws_o:
            obs.append('Sin uso en el programa')
            if s:
                aviso('Sin uso', 'Marca declarada sin uso', o, f'{s.get("Symbol")} — {s.get("Comment")}')
        elif ws_o and not reads_by.get(o) and not inv.get(o) and RE_BIT.match(o):
            obs.append('Se escribe pero no se lee en el programa (posible señal para HMI/SCADA)')
        if not s and (ws_o or uso):
            obs.append('AVISO: usada sin declarar en la tabla de símbolos')
            aviso('Aviso', 'Usada sin declarar', o, 'Marca sin entrada en la tabla de símbolos', uso)
        wb_ = sorted({f'{w["block"]} NW{w["nw"]}' for w in ws_o})
        rb_ = sorted({f'{r["block"]} NW{r["nw"]}' for r in reads_by.get(o, [])})
        r = wsm.max_row + 1
        wsm.append([None, s.get('Symbol', ''), s.get('DataType', '') or ('BOOL' if RE_BIT.match(o) else ''),
                    s.get('Comment', ''), traducir(s.get('Comment', '')), cut('\n'.join(act)), cut('\n'.join(des)),
                    cut('\n'.join(tgt[:40])), rem, cut(', '.join(wb_)), cut(', '.join(rb_)), '; '.join(obs)])
        link_addr(wsm, r, 1, o)
    finish(wsm, 12)

    # ===================================================== Temporizadores y contadores
    wst = sheet(wb, 'Temporizadores y contadores', ['Dirección / instancia', 'Símbolo', 'Comentario', 'Tipo',
                                                    'Preselección', 'Condición de arranque', 'Uso (dónde se consulta)',
                                                    'Bloque/Segmento', 'Observaciones'],
                [16, 22, 30, 30, 18, 70, 45, 45, 40], wrap_cols=(3, 4, 6, 7, 8, 9))
    TT = {'SE': 'SE — S_EVERZ retardo a la conexión (≈TON)', 'SA': 'SA — S_AVERZ retardo a la desconexión (≈TOF)',
          'SI': 'SI — S_IMPULS impulso (≈TP)', 'SV': 'SV — S_VIMP impulso prolongado',
          'SS': 'SS — S_SEVERZ retardo a la conexión con memoria', 'ZV': 'ZV — contador ascendente',
          'ZR': 'ZR — contador descendente'}
    tm_by = defaultdict(list)
    for t in TM:
        tm_by[t['op']].append(t)
    tz = sorted({o for o in SYM if RE_TZ.match(o)} | {o for o in used_ops if RE_TZ.match(o)} | set(tm_by),
                key=addr_key)
    for o in tz:
        s = SYM.get(o, {})
        evs = tm_by.get(o, [])
        tipos = sorted({TT.get(e['type'], e['type']) for e in evs})
        pres = sorted({str(e['preset']) for e in evs if e['preset']})
        conds = [f'• {render(e["cond"], True)} ({ubicacion(e)})' for e in evs]
        uso = sorted({f'{r["block"]} NW{r["nw"]}' for r in reads_by.get(o, [])})
        obs = []
        if len({(e['block'], e['nw'], tuple(c[0] for c in e['chain'])) for e in evs}) > 1:
            obs.append('AVISO: se arranca en varios sitios')
            aviso('Aviso', 'Temporizador arrancado en varios sitios', o, f'{len(evs)} arranques; preselecciones: {", ".join(pres)}',
                  '; '.join(sorted({ubicacion(e) for e in evs}))[:900])
        if len(pres) > 1:
            obs.append('Preselecciones distintas según el sitio de arranque')
        if not evs and not uso and s:
            obs.append('Sin uso en el programa')
            aviso('Sin uso', 'Temporizador/contador declarado sin uso', o, f'{s.get("Symbol")} — {s.get("Comment")}')
        if not evs and uso:
            obs.append('Se consulta pero no se arranca en el código analizado')
        r = wst.max_row + 1
        wst.append([None, s.get('Symbol', ''), s.get('Comment', ''), '\n'.join(tipos) or ('Contador' if o.startswith('Z') else 'S5'),
                    ', '.join(pres), cut('\n'.join(conds)), cut(', '.join(uso)),
                    cut('\n'.join(sorted({ubicacion(e) for e in evs}))), '; '.join(obs)])
        link_addr(wst, r, 1, o)
    # IEC (SFB) y temporizadores por software dentro de FB (INT en décimas)
    for c in D['CALLS']:
        if c['callee'] in ('SFB4', 'SFB5', 'SFB3', 'SFB0', 'SFB1', 'SFB2'):
            wst.append([f'{c["callee"]} {c["inst"] or ""}', '', '', c['callee'], c['params'].get('PT') or c['params'].get('PV') or '',
                        f'IN := {c["params"].get("IN") or c["params"].get("CU")}', '', f'{c["caller"]} NW{c["nw"]}', 'Temporizador IEC'])
    finish(wst, 9)

    # ===================================================== DB y analógicas
    wsd = sheet(wb, 'DB y analógicas', ['DB / variable', 'Símbolo DB', 'Variable', 'Tipo', 'Dirección', 'Valor inicial',
                                        'Comentario', 'Escalado', 'Rango', 'Uso'],
                [16, 24, 34, 14, 9, 12, 40, 26, 16, 50], wrap_cols=(7, 8, 10))
    # analógicas / periferia
    ana = sorted({x[0] for x in XR if re.match(r'^(PEW|PAW|PED|PAD|PEB|PAB|EW|AW|ED|AD)\d+', x[0])}, key=addr_key)
    wsd.append(['— PERIFERIA / PALABRAS DE E/S —'] + [''] * 9)
    wsd.cell(row=wsd.max_row, column=1).fill = F_SUB
    for o in ana:
        s = SYM.get(o, {})
        uso = bloques_uso(o)
        esc = 'No localizado (no hay FC105/FC106 ni escalado explícito)'
        if o.startswith(('PEW5', 'PAW5')):
            esc = 'Palabra de proceso PZD de variador (FB149 «Frequency Units» / MM420): consigna/estado, sin escalado analógico'
        elif o.startswith(('EW13', 'EW14', 'EW15')):
            esc = 'Lectura de báscula / palabra de entrada (FB40…FB90): ver bloque'
        r = wsd.max_row + 1
        wsd.append([None, '', s.get('Symbol', ''), s.get('DataType', '') or 'WORD', o, '', s.get('Comment', ''), esc,
                    'No localizado', uso])
        link_addr(wsd, r, 1, o)
    db_uso = defaultdict(set)
    for x in XR:
        m = re.match(r'^(DB\d+)', x[0])
        if m:
            db_uso[m.group(1)].add(f'{x[1]} NW{x[2]}')
    for c in D['CALLS']:
        if c['inst'] and re.match(r'^DB\d+$', c['inst']):
            db_uso[c['inst']].add(f'{c["caller"]} NW{c["nw"]} (instancia de {c["callee"]})')
    dbs = sorted([b for b in blocks.values() if b['Type'] == 'DB'], key=lambda b: int(b['Name'][2:]))
    nrows = 0
    for b in dbs:
        n = b['Name']
        tipo = f'DB de instancia de FB{b["FB"]}' if b.get('IsInstance') else 'DB global'
        wsd.append([f'— {n} —', b.get('Symbol') or '', '', tipo, '', '', b.get('Title') or '', '', '',
                    cut(', '.join(sorted(db_uso.get(n, [])))[:1500])])
        for c in range(1, 11):
            wsd.cell(row=wsd.max_row, column=c).fill = F_SUB
        for path, typ, addr, com, sv in D['dbvars'].get(n, [])[:400]:
            nrows += 1
            acc = sorted({f'{x[1]} NW{x[2]}' for x in XR if x[0] == f'{n}.{path}' or
                          D['abs2sym'].get(x[0]) == f'{n}.{path}'}) if not b.get('IsInstance') else []
            wsd.append([f'{n}.{path}', b.get('Symbol') or '', path, typ, addr, '' if sv in (None, 'None') else str(sv),
                        com, '', '', ', '.join(acc[:30])])
        if len(D['dbvars'].get(n, [])) > 400:
            wsd.append([f'{n}', '', f'… {len(D["dbvars"][n]) - 400} variables más (recortado)', '', '', '', '', '', '', ''])
    finish(wsd, 10)

    # ===================================================== Matriz causa-efecto
    cols = outputs_written
    rows_ = sorted({src for t in cols for src in deps.get(t, {}) if (src.startswith('E') or src.startswith('M'))
                    and RE_BIT.match(src)}, key=addr_key)
    wsc = wb.create_sheet('Matriz causa-efecto')
    wsc.cell(row=1, column=1, value='Causa \\ Efecto').fill = F_HEAD
    wsc.cell(row=1, column=1).font = FONT_HEAD
    wsc.cell(row=1, column=2, value='Símbolo').fill = F_HEAD
    wsc.cell(row=1, column=2).font = FONT_HEAD
    for j, t in enumerate(cols, 3):
        s = SYM.get(t, {})
        c = wsc.cell(row=1, column=j, value=f'{t} {s.get("Symbol", "")}')
        c.fill = F_HEAD
        c.font = FONT_HEAD
        c.alignment = Alignment(text_rotation=90, vertical='bottom', horizontal='center')
        wsc.column_dimensions[get_column_letter(j)].width = 4.2
    wsc.row_dimensions[1].height = 150
    wsc.column_dimensions['A'].width = 10
    wsc.column_dimensions['B'].width = 26
    colidx = {t: j for j, t in enumerate(cols, 3)}
    for i, src in enumerate(rows_, 2):
        wsc.cell(row=i, column=1, value=src)
        if src in anchor:
            wsc.cell(row=i, column=1).hyperlink = f"#'Referencias cruzadas'!A{anchor[src]}"
            wsc.cell(row=i, column=1).font = FONT_LINK
        wsc.cell(row=i, column=2, value=SYM.get(src, {}).get('Symbol', ''))
        for t, roles in inv.get(src, {}).items():
            if t in colidx:
                v = '/'.join(x for x in ('A', 'P', 'B') if x in roles)
                c = wsc.cell(row=i, column=colidx[t], value=v)
                c.alignment = Alignment(horizontal='center')
    wsc.freeze_panes = 'C2'
    last = f'{get_column_letter(len(cols) + 2)}{len(rows_) + 1}'
    rng = f'C2:{last}'
    wsc.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("B",C2))'], fill=F_B))
    wsc.conditional_formatting.add(rng, FormulaRule(formula=['C2="A"'], fill=F_A))
    wsc.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("A",C2))'], fill=F_A))
    wsc.conditional_formatting.add(rng, FormulaRule(formula=['C2="P"'], fill=F_P))
    wsc.auto_filter.ref = f'A1:{last}'


    # ===================================================== Dependencias + Consulta
    def origenes(dst_, maxd=6):
        """Entradas que influyen en dst_: directo o a través de marcas/DB/temporizadores."""
        res = {}
        frontier = [(dst_, 0, '')]
        seen = {dst_}
        while frontier:
            n, d, via = frontier.pop(0)
            for src, roles in deps.get(n, {}).items():
                if src in seen or src.startswith('~') or src.startswith('?'):
                    continue
                seen.add(src)
                if src.startswith('E') and RE_BIT.match(src):
                    res[src] = (d + 1, roles, via)
                elif d + 1 < maxd and not (src.startswith('A') and RE_BIT.match(src)):
                    frontier.append((src, d + 1, via or src))
        return res

    ROL = {'A': 'Activa', 'P': 'Permite', 'B': 'Bloquea'}
    wsdep = sheet(wb, 'Dependencias', ['Salida', 'Símbolo salida', 'Descripción salida', 'Entrada', 'Símbolo entrada',
                                       'Descripción entrada', 'Rol', 'Relación', 'A través de', 'Clave búsqueda',
                                       'Fila coincidente'],
                  [10, 18, 34, 10, 18, 34, 18, 12, 22, 10, 10], wrap_cols=(3, 6))
    dep_rows = 0
    for t in outputs_written:
        st = SYM.get(t, {})
        for src, (d, roles, via) in sorted(origenes(t).items(), key=lambda kv: (kv[1][0], addr_key(kv[0]))):
            ss = SYM.get(src, {})
            rol = ' / '.join(ROL[x] for x in ('A', 'P', 'B') if x in roles) if d == 1 else \
                ('Bloquea (indirecto)' if roles == {'B'} else 'Condiciona (indirecto)')
            r = wsdep.max_row + 1
            wsdep.append([t, st.get('Symbol', ''), traducir(st.get('Comment', '')) or st.get('Comment', ''), src,
                          ss.get('Symbol', ''), traducir(ss.get('Comment', '')) or ss.get('Comment', ''), rol,
                          'Directa' if d == 1 else f'Indirecta ({d} niveles)', nombre(via) if via else '—',
                          f'=A{r}&" "&B{r}&" "&C{r}&" "&D{r}&" "&E{r}&" "&F{r}',
                          f'=IF(Consulta!$C$4="","",IF(ISNUMBER(SEARCH(Consulta!$C$4,'
                          f'IF(Consulta!$C$5="Entrada",D{r}&" "&E{r}&" "&F{r},A{r}&" "&B{r}&" "&C{r}))),ROW(),""))'])
            dep_rows += 1
    finish(wsdep, 11)
    wsdep.column_dimensions['J'].hidden = True

    wsq = wb.create_sheet('Consulta', 1)
    wsq.column_dimensions['A'].width = 3
    for col, w_ in zip('BCDEFGHIJ', (10, 18, 34, 10, 18, 40, 22, 16, 22)):
        wsq.column_dimensions[col].width = w_
    wsq['B2'] = 'Consulta rápida: ¿qué entradas activan / permiten / bloquean una salida?'
    wsq['B2'].font = Font(bold=True, size=14, color=C_HEAD)
    wsq['B4'] = 'Buscar:'
    wsq['B4'].font = Font(bold=True)
    wsq['C4'] = 'A22.3'
    wsq['C4'].fill = PatternFill('solid', fgColor='FFF7D6')
    wsq['C4'].font = Font(bold=True, size=12)
    wsq['E4'] = ('Escriba una dirección (A22.3), un símbolo (422KM22.3), un BMK de motor (46M2) o una palabra del '
                 'comentario (Abschieber, Linie 603…). No distingue mayúsculas.')
    wsq.merge_cells('E4:J4')
    wsq['E4'].alignment = WRAP
    wsq.row_dimensions[4].height = 32
    wsq['B5'] = 'Buscar en:'
    wsq['B5'].font = Font(bold=True)
    wsq['C5'] = 'Salida'
    wsq['C5'].fill = PatternFill('solid', fgColor='FFF7D6')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Salida,Entrada"', allow_blank=False)
    wsq.add_data_validation(dv)
    dv.add('C5')
    wsq['E5'] = ('«Salida»: muestra las entradas que influyen en las salidas que coinciden (p. ej. motor en fallo → '
                 'qué entradas lo activan). «Entrada»: muestra las salidas afectadas por las entradas que coinciden.')
    wsq.merge_cells('E5:J5')
    wsq['E5'].alignment = WRAP
    wsq.row_dimensions[5].height = 32
    wsq['B6'] = '=COUNT(Dependencias!K:K)&" coincidencias (se muestran hasta 400)"'
    wsq['B6'].font = Font(italic=True, color='555555')
    hdr = ['Salida', 'Símbolo salida', 'Descripción salida', 'Entrada', 'Símbolo entrada', 'Descripción entrada',
           'Rol', 'Relación', 'A través de']
    for j, h in enumerate(hdr, 2):
        c = wsq.cell(row=8, column=j, value=h)
        c.fill = F_HEAD
        c.font = FONT_HEAD
    src_cols = 'ABCDEFGHI'
    for i in range(400):
        rr = 9 + i
        for j, sc in enumerate(src_cols, 2):
            c = wsq.cell(row=rr, column=j,
                         value=f'=IFERROR(INDEX(Dependencias!{sc}:{sc},SMALL(Dependencias!$K:$K,{i + 1})),"")')
            c.alignment = WRAP if j in (4, 7, 10) else TOP
    rng = 'B9:J408'
    wsq.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("Bloquea",$H9))'], fill=F_B))
    wsq.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("Activa",$H9))'], fill=F_A))
    wsq.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("Permite",$H9))'], fill=F_P))
    wsq.freeze_panes = 'B9'
    wsq['B410'] = ('Fuente: hoja «Dependencias» (todas las relaciones entrada → salida, directas e indirectas a través '
                   'de marcas, DB y temporizadores, hasta 6 niveles). Para el detalle de la condición, siga el enlace de '
                   'la salida en la hoja «Salidas».')

    # ===================================================== símbolos sin uso / sin declarar (otros)
    for o, s in SYM.items():
        if re.match(r'^(FB|FC)\d+$', o) and o not in blocks:
            aviso('Discrepancia', 'Símbolo de bloque inexistente', o, f'{s["Symbol"]}: el bloque no existe en el programa')
        m = re.match(r'^(\d{3})(KM|[A-Z])(\d+)\.(\d)$', s['Symbol'])
        mo = RE_BIT.match(o)
        if m and mo and f'{m.group(3)}.{m.group(4)}' != f'{mo.group(2)}.{mo.group(3)}':
            aviso('Discrepancia', 'Símbolo con dirección distinta', o,
                  f'El símbolo {s["Symbol"]} contiene la dirección {m.group(3)}.{m.group(4)} pero está asignado a {o}')
        m = re.match(r'^([EA])\s?(\d+)\.(\d)$', s['Symbol'])
        if m and mo and (m.group(1) != mo.group(1) or f'{m.group(2)}.{m.group(3)}' != f'{mo.group(2)}.{mo.group(3)}'):
            aviso('Discrepancia', 'Símbolo con dirección distinta', o,
                  f'El símbolo «{s["Symbol"]}» parece otra dirección ({m.group(1)}{m.group(2)}.{m.group(3)}) y está asignado a {o}')
    for n in D['no_llamados']:
        b = blocks[n]
        aviso('Info', 'Bloque no llamado', n, f'{b.get("Symbol") or ""} — {b.get("Title") or ""}: no se llama desde ningún OB')
    ind_blocks = sorted({e[2] for t in [D['tpl']] for n, tp in t.items() for e in [(None, None, n)] if tp['indirect']})
    for n in ind_blocks:
        aviso('Info', 'Direccionamiento indirecto', n, 'El bloque usa punteros, listas de salto (SPL) o accesos indexados: '
              'las referencias cruzadas de las zonas indirectas pueden estar incompletas')
    aviso('Info', 'Plano no disponible', '—',
          'No se ha podido contrastar con el esquema eléctrico ni el mecánico/neumático (no adjuntos en esta sesión). '
          'Las hojas de plano se han deducido del BMK incluido en el símbolo (convención Langhammer hoja = 200+byte para '
          'entradas y 400+byte para salidas) y deben verificarse.')

    # ===================================================== Avisos y discrepancias
    wsa = sheet(wb, 'Avisos y discrepancias', ['Severidad', 'Tipo', 'Variable / bloque', 'Símbolo', 'Detalle', 'Ubicación'],
                [13, 30, 16, 24, 80, 70], wrap_cols=(5, 6))
    sev_ord = {'Aviso': 0, 'Discrepancia': 1, 'Info': 2, 'Sin uso': 3}
    seen = set()
    for a in sorted(avisos, key=lambda a: (sev_ord.get(a[0], 9), a[1], addr_key(a[2]))):
        if a in seen:
            continue
        seen.add(a)
        r = wsa.max_row + 1
        wsa.append([a[0], a[1], None, SYM.get(a[2], {}).get('Symbol', ''), a[3], cut(a[4], 1500)])
        link_addr(wsa, r, 3, a[2])
    finish(wsa, 6)
    n = wsa.max_row
    wsa.conditional_formatting.add(f'A2:F{n}', FormulaRule(formula=['$A2="Aviso"'], fill=F_RED))
    wsa.conditional_formatting.add(f'A2:F{n}', FormulaRule(formula=['$A2="Discrepancia"'], fill=F_ORANGE))
    wsa.conditional_formatting.add(f'A2:F{n}', FormulaRule(formula=['$A2="Sin uso"'], fill=F_GRAY))
    # formato condicional en hojas de E/S y marcas
    for ws, col in ((wse, 'P'), (wss, 'R'), (wsm, 'L'), (wst, 'I')):
        m = ws.max_row
        ref = f'A2:{col}{m}'
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'ISNUMBER(SEARCH("AVISO",${col}2))'], fill=F_RED))
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'ISNUMBER(SEARCH("DISCREPANCIA",${col}2))'], fill=F_ORANGE))
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'ISNUMBER(SEARCH("Sin uso",${col}2))'], fill=F_GRAY))
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'ISNUMBER(SEARCH("Sin escritura",${col}2))'], fill=F_GRAY))

    # ===================================================== Portada
    cnt = defaultdict(int)
    for a in set(avisos):
        cnt[a[0]] += 1
    tb = defaultdict(int)
    for b in blocks.values():
        tb[b['Type']] += 1
    wr_out = len(outputs_written)
    ws0.column_dimensions['A'].width = 3
    ws0.column_dimensions['B'].width = 38
    ws0.column_dimensions['C'].width = 90
    ws0['B2'] = f'Documentación E/S y causa-efecto — {CFG["maquina"]}'
    ws0['B2'].font = Font(bold=True, size=16, color=C_HEAD)
    rows0 = [
        ('Máquina / línea', CFG['maquina']),
        ('Cliente', CFG.get('cliente', 'No indicado')),
        ('Proyecto STEP 7', CFG['proyecto']),
        ('Equipo / CPU', CFG['cpu']),
        ('Programa', CFG['programa']),
        ('Versión del programa', CFG['version']),
        ('Software de ingeniería', CFG['software']),
        ('Fecha del análisis', CFG['fecha']),
        ('Archivos analizados', CFG['archivos']),
        ('Archivos NO disponibles', CFG['faltan']),
        ('', ''),
        ('PLANO ELÉCTRICO (configurable)', ''),
        ('Nombre del PDF del plano', CFG.get('pdf', 'plano_electrico.pdf')),
        ('Desfase página PDF − n.º de hoja', 0),
        ('', 'Los enlaces «Página plano» abren <PDF>#page=<hoja + desfase>. Coloque el Excel y el PDF en la misma '
             'carpeta, escriba arriba el nombre exacto del PDF y, si la hoja 200 del plano no es la página 200 del PDF, '
             'ajuste el desfase. Nota: Excel abre el PDF con el visor predeterminado; algunos visores ignoran el '
             'parámetro #page.'),
        ('', ''),
        ('RESUMEN', ''),
        ('Bloques', ', '.join(f'{k}: {v}' for k, v in sorted(tb.items(), key=lambda t: order.get(t[0], 9)))),
        ('Entradas (bits E)', f'{stats["E"]} — con hoja de plano deducida del BMK: {stats["E_plano"]} '
                              f'({100 * stats["E_plano"] / max(stats["E"], 1):.0f} %)'),
        ('Salidas (bits A)', f'{stats["A"]} — escritas en el programa: {wr_out}; con hoja deducida del BMK: '
                             f'{stats["A_plano"]} ({100 * stats["A_plano"] / max(stats["A"], 1):.0f} %)'),
        ('Marcas', f'{len(marks)}'),
        ('Temporizadores / contadores', f'{len(tz)}'),
        ('Símbolos en la tabla', f'{len(SYM)}'),
        ('Avisos', f'{cnt["Aviso"]} avisos, {cnt["Discrepancia"]} discrepancias, {cnt["Sin uso"]} sin uso, {cnt["Info"]} informativos'),
        ('Contraste con plano eléctrico', '0 % verificado (plano no adjunto). Las columnas de plano indican «No localizado».'),
        ('', ''),
        ('CÓMO LEER ESTE LIBRO', ''),
        ('Método', 'Los bloques MC7 del archivado se han descompilado a AWL y se ha reconstruido la lógica booleana de cada '
                   '=, S y R (suma de productos, paréntesis, saltos SPB/SPBN/SPA/SPL). Las llamadas a FB/FC se han '
                   'expandido sustituyendo cada parámetro por el operando real del CALL, de modo que la condición de una '
                   'salida escrita dentro de un FB aparece con las entradas reales. Las variables estáticas de instancia '
                   'se muestran como DBn.variable.'),
        ('Notación booleana', '∧ = Y, ∨ = O, ¬ = NO, FP/FN = flanco, ⟨…⟩ = condición no booleana (comparación de '
                              'acumulador, lista de saltos), [a op b] = comparación. «SOLO SI» / «camino» = condición de '
                              'salto que debe cumplirse para que la instrucción se ejecute.'),
        ('Matriz causa-efecto', 'A = activa (primer contacto de una rama de activación), P = permite (contacto en serie '
                                'en la activación o condición de camino), B = bloquea (contacto negado en la activación '
                                'o contacto en la condición de RESET). Relación directa; las relaciones indirectas a '
                                'través de marcas se listan en la hoja Entradas.'),
        ('Colores', 'Rojo = aviso, naranja = discrepancia, gris = sin uso.'),
        ('Limitaciones', 'Temporizadores por software (contadores INT en décimas de segundo dentro de los FB '
                         'Langhammer) y accesos indirectos (punteros, DB[...]) no se resuelven por completo. La '
                         'remanencia de marcas es la de fábrica de la CPU (MB0–MB15) y no se ha verificado. Las '
                         'traducciones al español son orientativas (glosario).'),
    ]
    r = 4
    for k, v in rows0:
        ws0.cell(row=r, column=2, value=k).font = Font(bold=bool(k) and k.isupper(), color=C_HEAD if k.isupper() else '000000')
        c = ws0.cell(row=r, column=3, value=v)
        c.alignment = WRAP
        if k == 'Nombre del PDF del plano':
            assert CFG['_cell_pdf'] == r
            c.fill = PatternFill('solid', fgColor='FFF7D6')
        if k.startswith('Desfase'):
            assert CFG['_cell_off'] == r
            c.fill = PatternFill('solid', fgColor='FFF7D6')
        r += 1
    assert CFG['_cell_pdf'] and CFG['_cell_off']
    wb.move_sheet('Referencias cruzadas', offset=wb.sheetnames.index('Avisos y discrepancias') -
                  wb.sheetnames.index('Referencias cruzadas'))
    wb.save(dst)
    print('OK', dst, {k: v for k, v in stats.items()}, dict(cnt))


if __name__ == '__main__':
    main(*sys.argv[1:4])
